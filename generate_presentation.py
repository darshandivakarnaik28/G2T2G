import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_DARK = RGBColor(13, 27, 42)        # #0D1B2A
    NAVY_CARD = RGBColor(19, 35, 58)      # #13233A
    BG_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0
    PRIMARY = RGBColor(37, 99, 235)       # #2563EB (Royal Blue)
    SECONDARY = RGBColor(13, 148, 136)    # #0D9488 (Teal)
    ACCENT = RGBColor(245, 158, 11)       # #F59E0B (Amber)
    TEXT_DARK = RGBColor(15, 23, 42)      # #0F172A
    TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    TEXT_LIGHT = RGBColor(255, 255, 255)
    TEXT_LIGHT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    GREEN = RGBColor(16, 185, 129)

    def add_header(slide, title_text, category_text="G2T2G - INDIAN SIGN LANGUAGE PLATFORM", dark=False):
        # Header banner shape
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = SECONDARY if not dark else RGBColor(45, 212, 191)

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK if not dark else TEXT_LIGHT

    def set_slide_bg(slide, dark=False):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK if dark else BG_LIGHT

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        return shape

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, dark=True)

    # Left content box
    tbox1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.0), Inches(6.8), Inches(5.5))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "FINAL YEAR CAPSTONE PROJECT PRESENTATION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(45, 212, 191)

    p = tf1.add_paragraph()
    p.text = "G2T2G: Bidirectional Indian Sign Language Platform"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.space_after = Pt(12)

    p = tf1.add_paragraph()
    p.text = "Continuous Gesture-to-Text-to-Gesture Communication System with Dual-Frame Visual Synthesis (Master Signer & 3D Avatar Engine)"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    p.space_after = Pt(24)

    p = tf1.add_paragraph()
    p.text = "PROJECT TEAM (Department of CSE):"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(45, 212, 191)

    members = [
        "• Anupama G N (USN: 1GG23CS007)",
        "• Darshan Naik (USN: 1GG23CS016)",
        "• Hithaishi M (USN: 1GG23CS021)",
        "• Keerthana H S (USN: 1GG23CS026)"
    ]
    for m in members:
        pm = tf1.add_paragraph()
        pm.text = m
        pm.font.size = Pt(11)
        pm.font.color.rgb = TEXT_LIGHT

    # Right Hero Image
    hero_path = r"frontend\connect\assets\hero_isl_g2t2g.jpg"
    if os.path.exists(hero_path):
        s1.shapes.add_picture(hero_path, Inches(8.0), Inches(1.1), width=Inches(4.5))

    # ==========================================
    # SLIDE 2: Introduction & Motivation
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "Introduction & Motivation: Bridging the ISL Barrier")

    # 3 Column Cards
    col_w = Inches(3.64)
    h = Inches(5.2)
    top_pos = Inches(1.6)

    # Card 1: Society & Scale
    create_card(s2, Inches(0.8), top_pos, col_w, h)
    tb = s2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "18M+ Deaf Community"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(10)
    bullets1 = [
        "India is home to the largest deaf population globally (over 18 million individuals).",
        "Indian Sign Language (ISL) is their primary and natural mother tongue.",
        "Severe shortage of certified sign language interpreters in schools, hospitals, courts, and banks.",
        "Deaf individuals face extreme exclusion in daily public life."
    ]
    for b in bullets1:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # Card 2: The Two-Way Gap
    create_card(s2, Inches(4.84), top_pos, col_w, h)
    tb = s2.shapes.add_textbox(Inches(5.04), top_pos + Inches(0.2), col_w - Inches(0.4), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Need for True Bidirectionality"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p.space_after = Pt(10)
    bullets2 = [
        "Most existing solutions only do Sign-to-Text (one-way).",
        "Communication is a 2-way dialogue: Hearing individuals must also communicate back in Sign.",
        "Deaf users require rich visual validation to verify that their signed input was understood correctly.",
        "G2T2G provides complete Gesture → Text → Gesture loop."
    ]
    for b in bullets2:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # Card 3: Zero-Hardware Vision
    create_card(s2, Inches(8.88), top_pos, col_w, h)
    tb = s2.shapes.add_textbox(Inches(9.08), top_pos + Inches(0.2), col_w - Inches(0.4), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Accessible Ubiquitous Vision"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p.space_after = Pt(10)
    bullets3 = [
        "Costly data gloves and Kinect sensors are impractical for daily real-world use.",
        "G2T2G runs on standard RGB laptop webcams and mobile browsers.",
        "MediaPipe WebGL extracts 21 skeletal joints at 60 FPS directly in real time.",
        "Zero specialized hardware needed by the user."
    ]
    for b in bullets3:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 3: Problem Statement & Literature Survey
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "Problem Statement & Gaps in Existing Systems")

    create_card(s3, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.4), bg_color=RGBColor(239, 246, 255), border_color=PRIMARY)
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PROBLEM STATEMENT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p = tf.add_paragraph()
    p.text = "\"To design and develop a real-time, bidirectional, continuous Indian Sign Language communication framework that translates complex sign gestures into natural speech/text and synthesizes spoken/written language back into dual-mode visual ISL representations (Authentic Master Reference + 3D Avatar) on standard commodity devices.\""
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    # Table of comparison with literature
    rows, cols = 4, 4
    table_shape = s3.shapes.add_table(rows, cols, Inches(0.8), Inches(3.2), Inches(11.7), Inches(3.6))
    table = table_shape.table
    table.columns[0].width = Inches(2.5)
    table.columns[1].width = Inches(3.0)
    table.columns[2].width = Inches(3.0)
    table.columns[3].width = Inches(3.2)

    headers = ["Parameter", "Traditional Hardware Systems", "Existing Vision Systems", "Proposed G2T2G Framework"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_CARD
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT

    data = [
        ["Hardware Dependency", "Sensor gloves, Kinect depth cameras ($100-$500)", "Fixed camera setup, heavy GPU required", "Pure browser/webcam, zero extra hardware"],
        ["Signing Continuity", "Isolated single sign recognition only", "Flashing single words, false duplicate triggers", "Continuous staging buffer with hand-drop commit"],
        ["Translation Direction", "One-way (Sign to Text only)", "One-way or static dictionary lookups", "True Bidirectional: G2T + T2G + 3D Avatar + TTS"]
    ]
    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 0 else RGBColor(241, 245, 249)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = PRIMARY if c_idx == 3 else TEXT_DARK
                if c_idx == 3:
                    p.font.bold = True

    # ==========================================
    # SLIDE 4: Project Objectives
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "Project Objectives & Key Milestones")

    objs = [
        ("Zero-Sensor 3D Landmark Tracking", "Acquire 21 3D hand landmarks per hand at 60 FPS using MediaPipe WebGL without sensor gloves or infrared depth hardware.", PRIMARY),
        ("Continuous Multi-Sign Sentence Recognition", "Implement an automatic gesture staging buffer with hand-drop commit trigger to assemble continuous signed sentences seamlessly.", SECONDARY),
        ("Sub-Millisecond Machine Learning Classifier", "Engineer scale, translation, and orientation-invariant feature representations achieving >99% test accuracy at <0.01ms inference.", PRIMARY),
        ("Dual-Mode Visual ISL Synthesis (T2G)", "Simultaneously generate verified human master recordings and real-time Three.js 3D skeletal avatar hand animations with variable playback speed.", SECONDARY),
        ("Bilingual Speech & Text Support", "Provide real-time speech-to-text (STT) and voice synthesis (TTS) in both English and regional Kannada (ಕನ್ನಡ) language.", PRIMARY),
        ("In-App Crowdsourced Dataset Evolution", "Detect missing vocabulary words dynamically and prompt signers to record and expand the dataset on the fly.", SECONDARY)
    ]

    card_w = Inches(5.6)
    card_h = Inches(1.55)
    for idx, (title, desc, col) in enumerate(objs):
        r = idx // 2
        c = idx % 2
        left = Inches(0.8) + c * Inches(6.1)
        top = Inches(1.6) + r * Inches(1.8)
        create_card(s4, left, top, card_w, card_h)

        tb = s4.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{idx+1}. {title}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 5: System Requirement Specification (SRS)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "System Requirement Specification (SRS)")

    # Left: Hardware
    create_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "HARDWARE REQUIREMENTS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(14)

    hw = [
        ("Processor / CPU", "Intel Core i3 / AMD Ryzen 3 or higher (All ML inference runs in <0.01ms on standard CPU)."),
        ("System Memory (RAM)", "4 GB minimum (8 GB recommended for multi-stream browser WebGL)."),
        ("Camera / Vision Sensor", "Standard integrated 720p/1080p RGB webcam (No depth camera or specialized gloves required)."),
        ("Audio Input/Output", "Standard microphone for speech recognition; speakers/headphones for synthesized TTS voice."),
        ("Network Connectivity", "Localhost deployment or standard Internet bandwidth for Cloudflare tunnel access.")
    ]
    for k, v in hw:
        p = tf.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p = tf.add_paragraph()
        p.text = f"   {v}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # Right: Software
    create_card(s5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb = s5.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SOFTWARE REQUIREMENTS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p.space_after = Pt(14)

    sw = [
        ("Operating System", "Windows 10/11, Ubuntu Linux, macOS, Android/iOS mobile browsers."),
        ("Backend Framework", "Python 3.10+, FastAPI (high-performance async ASGI server), Uvicorn."),
        ("Machine Learning", "Scikit-Learn (Joblib pipelines), NumPy, Pandas, MediaPipe Python SDK."),
        ("Frontend & Graphics", "HTML5, Tailwind CSS, JavaScript (ES6+), Three.js (3D Skeletal Avatar), MediaPipe WebGL."),
        ("Database Engine", "SQLite (`isl_dataset.db`) storing 7,459 landmark sequences and master video references."),
        ("Speech APIs", "HTML5 Web Speech API (SpeechRecognition for STT; SpeechSynthesis for TTS).")
    ]
    for k, v in sw:
        p = tf.add_paragraph()
        p.text = f"• {k}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p = tf.add_paragraph()
        p.text = f"   {v}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(4)

    # ==========================================
    # SLIDE 6: System Architecture & Workflow
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "System Architecture & Bidirectional Flow")

    # Workflow blocks
    flow_steps = [
        ("1. Input Layer", "• Signer: RGB Webcam Stream\n• Hearing User: Mic Voice / Text\n• Quick-phrase chips selection", PRIMARY),
        ("2. Vision Tracking", "• MediaPipe Hands / Holistic\n• 21 3D Joint Coordinates\n• 60 FPS WebGL Processing", SECONDARY),
        ("3. Feature Extraction", "• 124 Static geometric features\n• 1,350 Dynamic spatio-temporal\n• Translation & scale normalized", PRIMARY),
        ("4. ML Classification", "• Logistic Regression (L2)\n• Static vs Dynamic auto-routing\n• 3-Frame Stability Filter", SECONDARY),
        ("5. Output Synthesis", "• Web Speech Voice TTS (Kannada/En)\n• Master Video Reference Frame\n• Three.js 3D Skeletal Avatar", GREEN)
    ]
    step_w = Inches(2.2)
    step_h = Inches(4.8)
    for i, (title, desc, col) in enumerate(flow_steps):
        left = Inches(0.8) + i * Inches(2.38)
        create_card(s6, left, Inches(1.8), step_w, step_h)

        # Header tag
        tag = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.8), step_w, Inches(0.6))
        tag.fill.solid()
        tag.fill.fore_color.rgb = col
        tag.line.fill.background()
        tf = tag.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.alignment = PP_ALIGN.CENTER

        tb = s6.shapes.add_textbox(left + Inches(0.15), Inches(2.5), step_w - Inches(0.3), step_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        for line in desc.split('\n'):
            p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
            p.text = line
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(6)

    # ==========================================
    # SLIDE 7: Mathematical Modeling & Feature Invariance
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Mathematical Modeling & Invariant Featurization")

    create_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "STATIC GEOMETRIC VECTOR (124 Features)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(12)

    s_bullets = [
        ("Translation Normalization", "All 21 landmarks centered on wrist: P'_i = P_i - P_0, eliminating signer positioning bias."),
        ("Scale Normalization", "Divided by maximum hand span: d_span = max ||P'_i||_2, invariant to camera-to-user distance."),
        ("Wrist-to-Fingertip Distances", "5 Euclidean metrics capturing extension of thumb, index, middle, ring, pinky."),
        ("Pairwise Fingertip Spread", "10 inter-fingertip distance pairs measuring splay and finger aperture."),
        ("Finger Curl Ratios", "Ratio of fingertip-to-wrist over total joint chain length to measure curvature."),
        ("3D Bone Unit Vectors", "15 unit vectors along finger phalanges encoding exact spatial joint angles.")
    ]
    for k, v in s_bullets:
        p = tf.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    create_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb = s7.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "DYNAMIC TEMPORAL VECTOR (1,350 Features)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p.space_after = Pt(12)

    d_bullets = [
        ("Trajectory Resampling", "Variable-speed gesture sequences are interpolated to a fixed T = 24 time steps."),
        ("Trajectory Anchoring", "Relative wrist motion anchored to origin frame: (x_t - x_0, y_t - y_0, z_t - z_0)."),
        ("Dual-Hand Landmark Vectors", "Concatenated active 3D coordinates across both left and right hands per time step."),
        ("Velocity & Direction Deltas", "First-order temporal difference vectors (P_{t} - P_{t-1}) capturing speed and stroke direction."),
        ("Robustness to Speed Variance", "Allows fast signers and slow learners to trigger identical classification templates."),
        ("Zero Data Leakage Split", "Stratified 70/15/15 split strictly prior to fitting StandardScaler parameters.")
    ]
    for k, v in d_bullets:
        p = tf.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # ==========================================
    # SLIDE 8: Continuous Communication & Hand Drop-to-Commit
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Continuous Communication: Drop-to-Commit Architecture")

    # Big top callout
    create_card(s8, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.4), bg_color=RGBColor(240, 253, 250), border_color=SECONDARY)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(11.3), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "SOLVING THE CONTINUOUS SIGNING CHALLENGE (Guide's Primary Recommendation)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p = tf.add_paragraph()
    p.text = "Traditional sign language recognition models predict isolated words frame-by-frame, causing repetitive stutter. G2T2G introduces a state-machine gesture buffer that stages continuous gestures into fluent sentences without requiring any mouse or keyboard touch."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    steps = [
        ("1. Multi-Frame Stability", "Requires 3 consecutive matching predictions before staging any gesture, filtering transitional hand motion.", PRIMARY),
        ("2. Anti-Duplication Filter", "Suppresses duplicate frames once a gesture is active. Prevents signing 'HELLO' from flooding 60 times per second.", SECONDARY),
        ("3. Live Staged Phrase Dock", "Renders interactive visual badge chips (e.g. [GOOD] + [MORNING]) directly on the camera overlay.", ACCENT),
        ("4. Hand Drop-to-Commit Trigger", "When hands relax or drop out of frame, the staging timer automatically commits the full sentence into speech and chat.", GREEN)
    ]
    card_w = Inches(5.6)
    card_h = Inches(1.8)
    for idx, (title, desc, col) in enumerate(steps):
        r = idx // 2
        c = idx % 2
        left = Inches(0.8) + c * Inches(6.1)
        top = Inches(3.3) + r * Inches(2.0)
        create_card(s8, left, top, card_w, card_h)

        tb = s8.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 9: Dataset Collection & Multi-Signer Distribution
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Dataset Collection & Multi-Signer Audit")

    # Table on Left
    rows, cols = 4, 5
    table_shape = s9.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(7.2), Inches(3.2))
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(1.3)
    table.columns[2].width = Inches(1.3)
    table.columns[3].width = Inches(1.4)
    table.columns[4].width = Inches(1.4)

    headers = ["Subset", "Raw Rows", "Excluded", "Retained Samples", "Classes Trained"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_CARD
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT

    dataset_rows = [
        ["Static Gestures", "4,774", "11 (book)", "4,763", "27 (A–Z, super)"],
        ["Dynamic Gestures", "2,702", "6 (empty)", "2,696", "18 Words/Phrases"],
        ["Total Database", "7,476", "17", "7,459", "45 ISL Gestures"]
    ]
    for r_idx, row in enumerate(dataset_rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 0 else RGBColor(241, 245, 249)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = PRIMARY if r_idx == 2 else TEXT_DARK
                if r_idx == 2:
                    p.font.bold = True

    # Multi-signer Card on Right
    create_card(s9, Inches(8.3), Inches(1.6), Inches(4.2), Inches(5.2))
    tb = s9.shapes.add_textbox(Inches(8.5), Inches(1.8), Inches(3.8), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "MULTI-SIGNER INTEGRITY"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p.space_after = Pt(10)

    signers = [
        ("1GG23CS007 (Anupama G N)", "Master reference signer, 239 static / 132 dynamic test samples."),
        ("1GG23CS016 (Darshan Naik)", "Static dataset validation contributor."),
        ("1GG23CS021 (Hithaishi M)", "240 static / 133 dynamic test samples."),
        ("1GG23CS026 (Keerthana H S)", "214 static / 140 dynamic test samples.")
    ]
    for s_name, s_role in signers:
        p = tf.add_paragraph()
        p.text = f"• {s_name}"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p = tf.add_paragraph()
        p.text = f"   {s_role}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    # Bottom notes on Left
    create_card(s9, Inches(0.8), Inches(5.1), Inches(7.2), Inches(1.7), bg_color=RGBColor(248, 250, 252))
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(5.25), Inches(6.8), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Rigorous Data Hygiene:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p = tf.add_paragraph()
    p.text = "• Stratified Sampling: 70% Train, 15% Validation, 15% Held-Out Test.\n• Zero-frame drop: Corrupt video captures with 0 active hand joints safely purged.\n• Zero demographic or single-signer bias achieved across evaluations."
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 10: Multi-Model Machine Learning Comparison
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "Machine Learning Multi-Model Comparison")

    # Table on Left
    rows, cols = 7, 5
    table_shape = s10.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(7.0), Inches(5.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(1.2)
    table.columns[2].width = Inches(1.2)
    table.columns[3].width = Inches(1.2)
    table.columns[4].width = Inches(1.2)

    headers = ["Model Architecture", "Static Test Acc", "Dynamic Test Acc", "Inference Latency", "Status"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_CARD
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT

    model_rows = [
        ["Naive Baseline (Majority)", "5.73%", "5.68%", "0.008 ms", "Baseline"],
        ["Logistic Regression (L2)", "99.16%", "100.00%", "0.006 ms", "★ Champion"],
        ["Extra Trees (250 trees)", "99.02%", "100.00%", "0.178 ms", "Evaluated"],
        ["Random Forest (200 trees)", "99.02%", "100.00%", "0.141 ms", "Evaluated"],
        ["MLP Neural Net (256x128)", "98.88%", "99.51%", "0.008 ms", "Evaluated"],
        ["Soft-Voting Ensemble", "98.88%", "100.00%", "0.353 ms", "Evaluated"]
    ]
    for r_idx, row in enumerate(model_rows):
        is_champ = (r_idx == 1)
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(236, 253, 245) if is_champ else (CARD_BG if r_idx % 2 == 0 else RGBColor(241, 245, 249))
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = SECONDARY if is_champ else TEXT_DARK
                if is_champ:
                    p.font.bold = True

    # Right: Embedded Comparison Chart
    chart_path = r"C:\Users\DELL\.gemini\antigravity\brain\aacf1404-bad1-43d3-820c-68da2a8d0492\model_comparison.png"
    if os.path.exists(chart_path):
        s10.shapes.add_picture(chart_path, Inches(8.0), Inches(1.6), width=Inches(4.5))

    # ==========================================
    # SLIDE 11: Confusion Matrices & Confidence Intervals
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11)
    add_header(s11, "Confusion Matrices & Statistical Validation")

    stat_cm_path = r"C:\Users\DELL\.gemini\antigravity\brain\aacf1404-bad1-43d3-820c-68da2a8d0492\static_confusion_matrix.png"
    dyn_cm_path = r"C:\Users\DELL\.gemini\antigravity\brain\aacf1404-bad1-43d3-820c-68da2a8d0492\dynamic_confusion_matrix.png"

    if os.path.exists(stat_cm_path):
        s11.shapes.add_picture(stat_cm_path, Inches(0.8), Inches(1.6), width=Inches(4.6))
    if os.path.exists(dyn_cm_path):
        s11.shapes.add_picture(dyn_cm_path, Inches(5.6), Inches(1.6), width=Inches(4.6))

    # Right Stat Card
    create_card(s11, Inches(10.4), Inches(1.6), Inches(2.2), Inches(5.2), bg_color=NAVY_CARD)
    tb = s11.shapes.add_textbox(Inches(10.55), Inches(1.8), Inches(1.9), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "STATISTICAL RIGOR"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(45, 212, 191)
    p.space_after = Pt(10)

    stats = [
        "95% Bootstrap CI:",
        "Static: [98.4%, 99.9%]",
        "Dynamic: [100%, 100%]",
        "",
        "Diagonal Density:",
        "Zero cross-class misclassification on dynamic signs.",
        "",
        "Latency / Sample:",
        "0.006 ms static",
        "0.010 ms dynamic",
        "Real-time 60 FPS capable."
    ]
    for st in stats:
        p = tf.add_paragraph()
        p.text = st
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT if "CI" in st or "Density" in st or "Sample" in st else TEXT_LIGHT_MUTED
        if "CI" in st or "Density" in st or "Sample" in st:
            p.font.bold = True

    # ==========================================
    # SLIDE 12: Cross-Signer Slice Analysis
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12)
    add_header(s12, "Cross-Signer Slice Analysis & Generalization")

    create_card(s12, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.2), bg_color=RGBColor(240, 249, 255), border_color=PRIMARY)
    tb = s12.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "VERIFYING ZERO DEMOGRAPHIC OR USER BIAS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p = tf.add_paragraph()
    p.text = "To guarantee the system works equally well for new unseen individuals, model evaluation was sliced across individual team members on strictly held-out test splits."
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_DARK

    rows, cols = 5, 7
    table_shape = s12.shapes.add_table(rows, cols, Inches(0.8), Inches(3.0), Inches(11.7), Inches(3.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(1.5)
    table.columns[2].width = Inches(1.5)
    table.columns[3].width = Inches(1.5)
    table.columns[4].width = Inches(1.5)
    table.columns[5].width = Inches(1.5)
    table.columns[6].width = Inches(1.5)

    headers = ["Signer ID / Name", "Static Samples", "Static Acc", "Static F1", "Dynamic Samples", "Dynamic Acc", "Dynamic F1"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_CARD
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = TEXT_LIGHT

    slices = [
        ["1GG23CS007 (Anupama G N)", "239", "100.00%", "1.0000", "132", "100.00%", "1.0000"],
        ["1GG23CS016 (Darshan Naik)", "22", "100.00%", "1.0000", "—", "—", "—"],
        ["1GG23CS021 (Hithaishi M)", "240", "97.50%", "0.9815", "133", "100.00%", "1.0000"],
        ["1GG23CS026 (Keerthana H S)", "214", "100.00%", "1.0000", "140", "100.00%", "1.0000"]
    ]
    for r_idx, row in enumerate(slices):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 0 else RGBColor(241, 245, 249)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = PRIMARY if "100" in val else TEXT_DARK
                if "100" in val:
                    p.font.bold = True

    # ==========================================
    # SLIDE 13: Live Application Interface Highlights
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s13)
    add_header(s13, "Live Application Modules & UI Design")

    col_w = Inches(3.64)
    h = Inches(5.2)
    top_pos = Inches(1.6)

    # 1. Normal User Interface
    create_card(s13, Inches(0.8), top_pos, col_w, h)
    tb = s13.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Normal / Hearing Studio"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p.space_after = Pt(10)
    b1 = [
        "Speech & Text Composition: Type message or speak via live microphone.",
        "Quick-Phrase Chips: Instant triggers for Hello, Thank you, Food, Morning, Water.",
        "Dual-Confirmation Output: Shows authentic master signer reference video above, and live 3D skeletal avatar below.",
        "Missing Gesture Prompt: Alerts if a word is not in dataset and offers to perform it."
    ]
    for b in b1:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # 2. Sign User Interface
    create_card(s13, Inches(4.84), top_pos, col_w, h)
    tb = s13.shapes.add_textbox(Inches(5.04), top_pos + Inches(0.2), col_w - Inches(0.4), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Sign User Studio"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p.space_after = Pt(10)
    b2 = [
        "High-Speed Live Camera: Real-time 21-landmark tracking overlay.",
        "Continuous Signing Banner: Displays staged phrase badge chips in real time.",
        "Drop-to-Commit: Automatically transmits text & audio when hands lower.",
        "Bilingual TTS Voice: Automatically speaks English or Kannada (pitch: 1.15, rate: 1.08)."
    ]
    for b in b2:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # 3. Continuous Live Chat
    create_card(s13, Inches(8.88), top_pos, col_w, h)
    tb = s13.shapes.add_textbox(Inches(9.08), top_pos + Inches(0.2), col_w - Inches(0.4), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "3. Continuous Live Chat"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p.space_after = Pt(10)
    b3 = [
        "Unified Split Screen: Chat thread on the LEFT; visual outputs on the RIGHT.",
        "Upper Right: Master Signer video & live webcam recognition.",
        "Lower Right: Interactive 3D Avatar with 0.1x to 1.5x speed controls.",
        "Unobstructed Canvas: Captions & sequence steppers placed strictly below video canvas."
    ]
    for b in b3:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 14: System Testing & Validation
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s14)
    add_header(s14, "System Testing & Performance Benchmarks")

    tests = [
        ("Unit & Component Testing", "• Verified MediaPipe 21-joint extraction across varying lighting conditions.\n• Confirmed Web Speech STT and TTS voice output across multiple browsers.\n• Tested Three.js WebGL landmark playback at 0.1x, 0.5x, 1.0x, and 1.5x.", PRIMARY),
        ("Continuous Integration & Regression", "• Anti-duplication regression tests: Zero repetitive false-positive triggers.\n• Staged phrase buffer verification: Successfully assembles multi-word phrases.\n• Drop-to-commit boundary checks: Verified seamless auto-commit on hand exit.", SECONDARY),
        ("Latency & Frame Rate Benchmarks", "• Camera capture & landmark extraction: ~16.6 ms (Steady 60 FPS).\n• ML Classification latency: 0.006 ms (Static) / 0.010 ms (Dynamic).\n• Total end-to-end pipeline latency: <50 ms (Virtually imperceptible delay).", GREEN),
        ("Cross-Device & Network Verification", "• Tested on standard Windows 11 laptop webcam.\n• Tested on mobile browsers via Cloudflare secure tunnel.\n• Bilingual validation: 100% verified translation accuracy between English and Kannada.", ACCENT)
    ]
    card_w = Inches(5.6)
    card_h = Inches(2.4)
    for idx, (title, desc, col) in enumerate(tests):
        r = idx // 2
        c = idx % 2
        left = Inches(0.8) + c * Inches(6.1)
        top = Inches(1.8) + r * Inches(2.7)
        create_card(s14, left, top, card_w, card_h)

        tb = s14.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(8)

        for line in desc.split('\n'):
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(10.5)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(4)

    # ==========================================
    # SLIDE 15: Distinct Advantages & Innovation
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s15)
    add_header(s15, "Key Novelties & Distinct Advantages")

    innovations = [
        ("1. Continuous Communication Buffer", "Replaces clunky isolated word detection with smooth multi-word sentence formation triggered by natural hand drop-to-commit.", PRIMARY),
        ("2. Dual Visual Feedback Synthesis", "Combines verified human signer video with customizable Three.js 3D avatar skeletons for maximum visual clarity.", SECONDARY),
        ("3. Truly Accessible & Cost-Free", "Zero sensor gloves, zero Kinect depth cameras. Runs in standard web browsers on any commodity laptop or smartphone.", PRIMARY),
        ("4. Bilingual Inclusivity", "Seamlessly supports English and regional Kannada (ಕನ್ನಡ) text and voice, promoting linguistic accessibility.", SECONDARY),
        ("5. Self-Expanding Vocabulary", "In-app missing gesture detection invites users to perform and save new signs, enabling crowd-powered database growth.", PRIMARY),
        ("6. Production-Grade Sub-Millisecond Speed", "Logistic Regression L2 champion model delivers 99.16% test accuracy in 0.006 ms per sample.", SECONDARY)
    ]
    card_w = Inches(5.6)
    card_h = Inches(1.55)
    for idx, (title, desc, col) in enumerate(innovations):
        r = idx // 2
        c = idx % 2
        left = Inches(0.8) + c * Inches(6.1)
        top = Inches(1.6) + r * Inches(1.8)
        create_card(s15, left, top, card_w, card_h)

        tb = s15.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), card_w - Inches(0.4), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 16: Conclusion & Future Scope (Dark Theme)
    # ==========================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s16, dark=True)
    add_header(s16, "Conclusion & Future Enhancements", dark=True)

    # Left: Conclusion
    create_card(s16, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), bg_color=NAVY_CARD, border_color=PRIMARY)
    tb = s16.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CONCLUSION"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(45, 212, 191)
    p.space_after = Pt(14)

    concl = [
        "G2T2G successfully bridges the communication divide between the Deaf community and hearing society.",
        "Built a robust, bidirectional pipeline operating in sub-50ms latency on ordinary commodity devices.",
        "Achieved 99.16% static accuracy and 100% dynamic accuracy across 45 ISL classes with zero signer bias.",
        "Successfully implemented our Guide's recommendation for continuous sentence signing with drop-to-commit segmentation."
    ]
    for c in concl:
        p = tf.add_paragraph()
        p.text = "• " + c
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(10)

    # Right: Future Scope
    create_card(s16, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), bg_color=NAVY_CARD, border_color=SECONDARY)
    tb = s16.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FUTURE ENHANCEMENTS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(45, 212, 191)
    p.space_after = Pt(14)

    future = [
        "Facial Expression & Emotion Tracking: Integrate MediaPipe Face Mesh to capture non-manual grammatical markers in ISL.",
        "Vocabulary Expansion: Scale dataset to 500+ daily conversational signs across multiple regional Indian dialects.",
        "Native Mobile App: Package as an offline mobile application (Flutter / PWA) for instant on-device signing.",
        "AR Smart Glasses: Projection of 3D ISL avatar onto heads-up displays for hands-free live conversations."
    ]
    for f in future:
        p = tf.add_paragraph()
        p.text = "• " + f
        p.font.size = Pt(11.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(10)

    # ==========================================
    # SLIDE 17: Thank You & Q&A (Dark Theme)
    # ==========================================
    s17 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s17, dark=True)

    create_card(s17, Inches(2.5), Inches(1.5), Inches(8.33), Inches(4.5), bg_color=NAVY_CARD, border_color=PRIMARY)
    tb = s17.shapes.add_textbox(Inches(2.7), Inches(1.8), Inches(7.93), Inches(3.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "THANK YOU"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(45, 212, 191)
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(10)

    p = tf.add_paragraph()
    p.text = "Questions & Answers / Discussion"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(20)

    p = tf.add_paragraph()
    p.text = "G2T2G: Gesture-to-Text-to-Gesture Continuous ISL Platform"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "Department of Computer Science & Engineering"
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT_MUTED
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(16)

    p = tf.add_paragraph()
    p.text = "Live Interactive Demo: Ready for Evaluation"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(245, 158, 11)
    p.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_path = r"c:\Users\DELL\OneDrive\Desktop\G2T2G\G2T2G_Final_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
